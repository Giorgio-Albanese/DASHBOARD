import io
import re
import traceback
import zipfile
import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter
import pandas as pd
import streamlit as st


def render_report_intermediari(conn):
  st.subheader("📊 Generatore Report per Intermediario")
  st.markdown(
      "Carica il file dei contatti (`Contatti.xlsx`) per elaborare e scaricare"
      " i report Excel suddivisi per intermediario e finanziaria."
  )

  if "zip_buffer_intermediari" not in st.session_state:
    st.session_state["zip_buffer_intermediari"] = None

  contatti_file = st.file_uploader(
      "Carica il file Contatti.xlsx", type=["xlsx"], key="uploader_contatti"
  )

  if contatti_file is not None:
    try:
      contatti_df = pd.read_excel(contatti_file)
      contatti_df = contatti_df.dropna(
          subset=["Finanziaria", "Intermediario"]
      ).copy()
      contatti_df["Finanziaria_Clean"] = (
          contatti_df["Finanziaria"]
          .astype(str)
          .str.upper()
          .str.replace(r"\s+", " ", regex=True)
          .str.strip()
      )
    except Exception as e:
      st.error(f"Errore nella lettura del file Contatti: {e}")
      return

    if st.button(
        "🚀 Genera Report Intermediari",
        type="primary",
        use_container_width=True,
    ):
      # Blocco protetto per intercettare qualsiasi errore e mostrarlo a schermo
      try:
        with st.spinner(
            "Estrazione dati da DuckDB ed elaborazione calcoli in corso..."
        ):
          # Estrazione dati da DuckDB
          tables = conn.execute("SHOW TABLES").fetchall()
          table_name = tables[0][0] if tables else "dati"
          df_db = conn.execute(f"SELECT * FROM {table_name}").df()

          # Controllo colonne obbligatorie per evitare KeyError silenziose
          colonne_richieste = [
              "DECORRENZA",
              "CONTRAENTE",
              "PREMI_NETTO",
              "PROVVACQ",
              "ESTINZIONI",
              "ESTINZIONI_PROVV",
              "LIQUIDAZIONI",
              "ID",
          ]
          colonne_mancanti = [
              col for col in colonne_richieste if col not in df_db.columns
          ]
          if colonne_mancanti:
            st.error(
                f"❌ Errore: Nel database mancano le colonne obbligatorie:"
                f" {colonne_mancanti}"
            )
            return

          df = df_db.copy()

          df["DECORRENZA_DT"] = pd.to_datetime(
              df["DECORRENZA"], format="mixed", errors="coerce"
          )
          df["Anno"] = df["DECORRENZA_DT"].dt.year
          df = df[df["Anno"].isin(range(2018, 2027))].copy()

          contraente_upper = (
              df["CONTRAENTE"]
              .astype(str)
              .str.upper()
              .str.replace(r"\s+", " ", regex=True)
              .str.strip()
          )

          finanziaria_pulita = np.select(
              [
                  contraente_upper.str.contains(r"I\.B\.L\.", regex=True),
                  contraente_upper.str.contains("CREDITIS", regex=True),
              ],
              ["IBL ISTITUTO BANCARIO DEL LAVORO", "CREDITIS"],
              default=contraente_upper,
          )

          df["FINANZIARIA_TEMP"] = finanziaria_pulita
          cleaned = df["FINANZIARIA_TEMP"].str.replace(
              r"[^\w\s]", " ", regex=True
          )
          cleaned = cleaned.str.replace(
              r"\b(SPA|SRL|S P A|S R L)\b", "", regex=True
          )
          df["FINANZIARIA"] = (
              cleaned.str.replace(r"\s+", " ", regex=True).str.strip()
          )

          def map_tipo_azienda(row):
            az = row.get("AZTIPO", None)
            prod = str(row.get("CODICEPROD", ""))
            if az == 0 or prod == "CQP":
              return "Pensionato"
            elif az in [1, 2]:
              return "Pubblico"
            elif az == 3:
              return "Parapubblico"
            elif az in [6, 7, 9, 10]:
              return "Medio Privato"
            elif az in [8, 11]:
              return "Grande Privato"
            elif az in [4, 5] or (az is not None and 12 <= az <= 20):
              return "Resto"
            elif az in [21, 22]:
              return "Del Pubblico"
            elif az == 23:
              return "Del Parapubblico"
            elif az == 26:
              return "Del Medio Privato"
            elif (az is not None and 24 <= az <= 25) or (
                az is not None and 40 <= az <= 50
            ):
              return "Del Resto"
            else:
              return "Privato"

          df["Tipo_Azienda"] = df.apply(map_tipo_azienda, axis=1)

          if "RAMO" in df.columns and "Tipo" not in df.columns:
            df["Tipo"] = df["RAMO"]
          elif "Tipo" not in df.columns:
            df["Tipo"] = "Generico"

          group_keys = ["Anno", "FINANZIARIA", "Tipo", "Tipo_Azienda"]
          df_grouped = (
              df.groupby(group_keys)
              .agg(
                  Premi_netto=("PREMI_NETTO", "sum"),
                  Provvigioni=("PROVVACQ", "sum"),
                  Estinzioni=("ESTINZIONI", lambda x: x.fillna(0).sum()),
                  Provvigioni_rec=("ESTINZIONI_PROVV", "sum"),
                  Sinistri=("LIQUIDAZIONI", "sum"),
                  N_Polizze=("ID", "count"),
                  Tipo_Azienda_prevalenza=(
                      "Tipo_Azienda",
                      lambda x: x.mode()[0] if not x.mode().empty else "Privato",
                  ),
              )
              .reset_index()
          )

          df_grouped["Premi_netto_est"] = (
              df_grouped["Premi_netto"]
              - df_grouped["Provvigioni"]
              + df_grouped["Estinzioni"]
              + df_grouped["Provvigioni_rec"]
          )
          df_grouped["Premi"] = np.maximum(1, df_grouped["Premi_netto_est"])

          df_grouped["Sinistri_attesi"] = (
              10
              / np.maximum(1, np.minimum(10, 2026 - df_grouped["Anno"]))
              * df_grouped["Sinistri"]
          )
          df_grouped["LR_netto"] = df_grouped["Sinistri"] / df_grouped["Premi"]
          df_grouped["LR"] = df_grouped["Sinistri"] / df_grouped["Premi_netto"]
          df_grouped["LR_netto_atteso"] = (
              df_grouped["Sinistri_attesi"] / df_grouped["Premi"]
          )
          df_grouped["LR_atteso"] = (
              df_grouped["Sinistri_attesi"] / df_grouped["Premi_netto"]
          )
          df_grouped["Provv_perc"] = (
              df_grouped["Provvigioni"] / df_grouped["Premi_netto"]
          )

          df_int = pd.merge(
              df_grouped,
              contatti_df[["Finanziaria_Clean", "Intermediario"]],
              left_on="FINANZIARIA",
              right_on="Finanziaria_Clean",
              how="left",
          )
          df_int["FINANZIARIA"] = df_int["Finanziaria"].fillna(
              df_int["FINANZIARIA"]
          )
          df_int = df_int.drop(
              columns=["Finanziaria_Clean", "Finanziaria"], errors="ignore"
          )
          cols = list(df_int.columns)
          if "Anno" in cols and "FINANZIARIA" in cols:
            cols.remove("FINANZIARIA")
            anno_idx = cols.index("Anno")
            cols.insert(anno_idx + 1, "FINANZIARIA")
            df_int = df_int[cols]

          # Creazione archivio ZIP in memoria
          zip_buffer = io.BytesIO()
          with zipfile.ZipFile(
              zip_buffer, "w", zipfile.ZIP_DEFLATED
          ) as zip_file:
            df_filtrato = df_int.dropna(subset=["Intermediario"])
            df_filtrato = df_filtrato[
                ~df_filtrato["Intermediario"].isin(["", " ", None])
            ]

            intermediari = df_filtrato["Intermediario"].unique()

            for int_nome in intermediari:
              int_pulito = re.sub(r'[\\/:*?"<>|]', "_", str(int_nome))
              df_sub_int = df_filtrato[df_filtrato["Intermediario"] == int_nome]

              fogli_lista = {
                  k: v
                  for k, v in df_sub_int.groupby("FINANZIARIA")
                  if k != "NA" and str(k) != "nan"
              }

              if not fogli_lista:
                continue

              excel_buffer = io.BytesIO()
              wb = openpyxl.Workbook()
              wb.remove(wb.active)

              for fin_nome, df_fin in fogli_lista.items():
                nome_sheet = re.sub(r'[\\/:*?"<>|]', "_", str(fin_nome))[:31]
                ws = wb.create_sheet(title=nome_sheet)

                headers = list(df_fin.columns)
                ws.append(headers)
                for row in df_fin.itertuples(index=False):
                  ws.append(list(row))

                ws.auto_filter.ref = ws.dimensions
                for col in ws.columns:
                  max_len = max(len(str(cell.value or "")) for cell in col)
                  col_letter = get_column_letter(col[0].column)
                  ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

              wb.save(excel_buffer)
              excel_buffer.seek(0)
              zip_file.writestr(f"{int_pulito}.xlsx", excel_buffer.read())

          zip_buffer.seek(0)
          st.session_state["zip_buffer_intermediari"] = zip_buffer
          st.success("Tutti i report per intermediario sono stati generati!")

      except Exception as e:
        # Se c'è un errore, lo mostra chiaramente a schermo invece di crashare in silenzio
        st.error(
            "❌ Si è verificato un errore imprevisto durante l'elaborazione:"
        )
        st.exception(e)
        traceback.print_exc()

  if st.session_state["zip_buffer_intermediari"] is not None:
    st.markdown("---")
    st.download_button(
        label="📦 Scarica Archivio ZIP (Report Intermediari)",
        data=st.session_state["zip_buffer_intermediari"],
        file_name="Report_Intermediari_Excel.zip",
        mime="application/zip",
        use_container_width=True,
    )
